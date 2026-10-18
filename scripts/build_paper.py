"""Rebuild the existing paper from the committed experiment tables.

Requires requirements-paper.txt. Run build_paper_figures.py first.
Export the resulting DOCX to PDF with a compatible office renderer.
"""

import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper/SMS_Spam_Detection_Paper.docx"


def read_csv(protocol, name):
    with (ROOT / "results" / protocol / name).open(newline="") as stream:
        return list(csv.DictReader(stream))


def read_json(protocol, name):
    return json.loads((ROOT / "results" / protocol / name).read_text())


def main():
    meta = read_json("full", "metadata.json")
    template = read_json("template", "metadata.json")
    clean = read_csv("full", "clean_metrics.csv")
    selected = meta["selected_model"]
    # The discussion is specific to these experiments; fail instead of silently
    # attaching old interpretations to a newly trained model or changed split.
    assert selected == "linear_svc_word_gentle"
    assert template["selected_model"] == "linear_svc_char_no_punctuation"
    assert meta["split_counts"]["test"] == {"ham": 904, "spam": 130, "total": 1034}
    tuned = next(r for r in clean if r["model"] == selected and r["threshold_mode"] == "validation")
    assert (int(tuned["fp"]), int(tuned["fn"])) == (4, 8)
    external = read_csv("external/imc25/evaluation", "metrics.csv")
    external_meta = read_json("external/imc25/evaluation", "metadata.json")
    external_check = read_json("external/imc25/evaluation", "verification.json")
    assert external_meta["model_id"] == selected
    assert external_meta["model_sha256"] == meta["models"][selected]["artifact_sha256"]
    assert external_check["status"] == "passed"
    assert external_check["output_hashes"] == external_meta["output_sha256"]
    for name, digest in external_meta["output_sha256"].items():
        path = ROOT / "results/external/imc25/evaluation" / name
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest

    def ext(cohort, mode="validation", view="published"):
        return next(r for r in external if r["cohort"] == cohort
                    and r["threshold_mode"] == mode and r["view"] == view)

    recent, broad = ext("recent_2023_2024"), ext("all_english")
    assert (int(recent["detected"]), int(recent["missed"]), int(recent["n_messages"])) == (176, 196, 372)
    assert (int(broad["detected"]), int(broad["n_messages"])) == (7491, 15531)
    primary_recall = f'{100*float(recent["recall"]):.2f}%'
    primary_interval = f'{100*float(recent["recall_lower"]):.2f}% to {100*float(recent["recall_upper"]):.2f}%'
    adapted = read_csv("adaptation/evaluation", "test_metrics.csv")
    adapted_meta = read_json("adaptation/evaluation", "evaluation_metadata.json")
    adapted_check = read_json("adaptation/evaluation", "verification.json")
    assert adapted_check["status"] == "passed"
    assert adapted_meta["promotion_gate_passed"] is False
    assert adapted_meta["selected_model"] == "adapted_word_char_c1"
    assert adapted_check["output_sha256"] == adapted_meta["output_sha256"]
    for name, digest in adapted_meta["output_sha256"].items():
        assert hashlib.sha256((ROOT / "results/adaptation/evaluation" / name).read_bytes()).hexdigest() == digest

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.75)
    section.footer_distance = Inches(0.3)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        for fonts in style.element.xpath(".//w:rFonts"):
            for attr in list(fonts.attrib):
                if "theme" in attr.lower():
                    del fonts.attrib[attr]
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    for name, size in [("Heading 1", 14), ("Heading 2", 12)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.bold = True
        doc.styles[name].paragraph_format.space_before = Pt(10)
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.size = Pt(9.5)
    doc.styles["Caption"].font.italic = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "SMS Spam Detection with Grouped and External Evaluation"
    doc.core_properties.author = "Kabiyev Altynbek"
    doc.core_properties.subject = "Reproducible evaluation of sparse SMS classifiers"
    doc.core_properties.comments = ""
    doc.core_properties.last_modified_by = "Kabiyev Altynbek"

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def h(text):
        doc.add_heading(text, level=1)

    def page():
        doc.add_page_break()

    def table(headers, rows, widths, caption):
        cap = p(caption, "Caption")
        cap.paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Inches(width)
        for cells, values, header in [(t.rows[0].cells, headers, True)] + [
                (t.add_row().cells, row, False) for row in rows]:
            for i, (cell, value, width) in enumerate(zip(cells, values, widths)):
                cell.width = Inches(width)
                cell.text = str(value)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    el = OxmlElement(f"w:{edge}")
                    for key, val in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                        el.set(qn(f"w:{key}"), val)
                    borders.append(el)
                tcpr.append(borders)
                margin = OxmlElement("w:tcMar")
                for edge, val in [("top", "65"), ("bottom", "65"), ("left", "80"), ("right", "80")]:
                    el = OxmlElement(f"w:{edge}")
                    el.set(qn("w:w"), val)
                    el.set(qn("w:type"), "dxa")
                    margin.append(el)
                tcpr.append(margin)
                if header:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EAEAEA")
                    tcpr.append(shade)
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_after = Pt(0)
                    para.paragraph_format.line_spacing = 1.05
                    for run in para.runs:
                        run.font.size = Pt(9.5)
                        run.bold = header
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        p("").paragraph_format.space_after = Pt(0)

    p(doc.core_properties.title, "Title")
    p("Kabiyev Altynbek\nAstana IT University, Kazakhstan")
    h("Abstract")
    p("This study evaluates eleven sparse classifiers on the historical English SMS Spam Collection, separating training search, validation thresholds and final testing. After deduplication, 5,171 messages remain in grouped splits. The selected word TF-IDF LinearSVC achieves spam F1 0.9531, with four false positives and eight misses among 1,034 test messages. Its conditional 95% group-bootstrap F1 interval is 0.9249 to 0.9769. Separator insertion lowers mean recall from 0.9385 to 0.6035. On 372 independent English reports with explicit 2023–2024 years, the unchanged model detects 176: recall " + primary_recall + ". We then refit with newer spam and ordinary financial messages under a separate frozen protocol. On another independent source, recall rises from 42.63% to 58.66%, but false positives on government notifications rise from 11/55 to 19/55. A separate financial holdout improves from F1 0.6400 to 0.8605. The adapted candidate fails the predeclared replacement criterion. Higher historical or spam-class F1 does not establish reliable transfer or acceptable harm to legitimate messages.")
    p("Keywords: SMS spam, TF-IDF, grouped evaluation, external evaluation, character obfuscation.")
    h("1 Introduction")
    p("SMS filtering must detect unwanted messages without suppressing legitimate communication. Class imbalance makes accuracy an incomplete measure: on this test set, always predicting ham achieves 87.43% accuracy but detects no spam. We therefore report spam precision, recall and F1 together with false-positive and false-negative counts. We also separate clean-text quality from sensitivity to character edits that can disrupt a vocabulary-based classifier.")
    p("The contribution is an auditable comparison of representations and decision rules. Candidates share training folds and attack realizations. Model choice is frozen before test scoring, and validation chooses the threshold. A second grouping rule examines residual template overlap. An independent collection of reported SMS then tests transfer with the original model and threshold. A separate adaptation follow-up uses those inspected reports for development and reserves new test data.")
    h("2 Related work")
    p("The SMS Spam Collection accompanies Almeida, Hidalgo and Yamakami's benchmark research [1, 2]. Its historical sources limit deployment claims. ExplainableDetector studies transformer detection and explainability [3]; SpamDam investigates newer SMS, federated learning and adversaries [4]; COPS studies an on-device smishing and URL pipeline [5]. Their metrics are not directly comparable to our splits. Agarwal et al. release public smishing reports collected during 2017–2024 [9]. We use that independent artifact for a separate frozen-model evaluation, without claiming to reproduce or outperform those systems.")

    page()
    h("3 Data and evaluation protocol")
    p("The bundled corpus contains 5,574 labeled messages [1]. Input checks reject missing or blank training text, unsupported labels and conflicting labels within a group. Removing 403 exact duplicate rows leaves 4,518 ham and 653 spam messages. The remaining original strings are preserved as model inputs. Lowercasing and collapsing whitespace define 5,159 split groups; punctuation and Unicode remain in these keys. Message and group hashes provide an auditable membership record.")
    table(["Partition", "Ham", "Spam", "Total", "Purpose"],
          [[s.title(), meta["split_counts"][s]["ham"], meta["split_counts"][s]["spam"], meta["split_counts"][s]["total"], role]
           for s, role in [("train", "Model search and fitting"), ("validation", "Threshold selection"), ("test", "Final reporting")]],
          [1.1, .7, .7, .7, 3.8], "Table 1. Default protocol partitions after exact deduplication.")
    p("Unique homogeneous groups are stratified by label into approximately 60% training, 20% validation and 20% test data, using seeds 42 and 43. Whole groups remain in one partition. Five shared StratifiedGroupKFold folds inside training use seed 42. Every cleaner, vocabulary and inverse-document-frequency transformation is fitted within its training fold. No validation or test text enters feature fitting. Grouping reduces a declared form of overlap; it does not prove sender, campaign or temporal independence.")
    h("4 Models and operating thresholds")
    p("The implementation uses scikit-learn pipelines [6]. Four original word-feature baselines are Multinomial Naive Bayes, Logistic Regression, linear-kernel SVC and Random Forest. Seven LinearSVC variants compare word and character representations under legacy cleanup, gentle cleanup, their word-plus-character combination, and character models with punctuation or currency symbols removed. Legacy cleanup lowercases text and retains only a-z, 0-9 and whitespace. Gentle cleanup lowercases and collapses whitespace without stripping symbols.")
    p("Word TF-IDF uses unigrams and bigrams; char_wb TF-IDF uses character 3- to 5-grams within word boundaries. Vocabulary budgets are 1,000 or 5,000 features; the combined representation divides the budget equally and normalizes the concatenation. Search uses alpha in {0.1, 1} for Naive Bayes; C in {0.5, 1} and class weights in {none, balanced} for linear models; and depth in {unlimited, 20} with minimum leaf size in {1, 2} for the 100-tree forest. There are 84 candidates and 420 fold fits.")
    p("Each variant retains its highest mean training-CV spam F1 configuration. The overall choice uses that same criterion, breaking model ties alphabetically. Best searched CV scores are selection scores, not independent estimates. Each chosen pipeline is refitted on training only. Its validation threshold maximizes recall subject to empirical FPR <= 1%, then prefers fewer false positives and a higher threshold. Decisions use score >= threshold. The model is not refitted afterward. Native default predictions are reported separately; the validation constraint does not guarantee future FPR [7].")

    page()
    h("5 Clean test results")
    names = {"nb_word_legacy": "NB word legacy", "lr_word_legacy": "LR word legacy", "svc_word_legacy": "SVC word legacy", "rf_word_legacy": "RF word legacy", "linear_svc_word_legacy": "LinearSVC word legacy", "linear_svc_char_legacy": "LinearSVC char legacy", "linear_svc_word_gentle": "LinearSVC word gentle*", "linear_svc_char_gentle": "LinearSVC char gentle", "linear_svc_word_char_gentle": "LinearSVC word + char", "linear_svc_char_no_punctuation": "LinearSVC char no punctuation", "linear_svc_char_no_currency": "LinearSVC char no currency"}
    rows = []
    for name in meta["models"]:
        default, val = [next(r for r in clean if r["model"] == name and r["threshold_mode"] == mode) for mode in ["default", "validation"]]
        rows.append([names[name], f'{meta["models"][name]["best_cv_f1"]:.4f}', f'{float(default["f1"]):.4f}', f'{float(val["f1"]):.4f}', f'{val["fp"]}/{val["fn"]}', f'{100*float(val["fpr"]):.2f}%'])
    table(["Variant", "CV F1", "Default F1", "Tuned F1", "FP/FN", "FPR"], rows,
          [2.65, .75, .95, .9, .85, .9], "Table 2. Default protocol. FP/FN and FPR use the validation threshold. *CV-selected model.")
    p("Training CV selects word/gentle LinearSVC (C = 1, balanced class weights, vocabulary limit 5,000). At threshold -0.08620419778818611, test precision is 0.9683 and recall is 0.9385. The confusion counts are TN = 900, FP = 4, FN = 8 and TP = 122. Accuracy is 0.9884, ROC-AUC is 0.9957 and average precision is 0.9817. Native default predictions yield three false positives and ten false negatives, so tuning exchanges one additional false positive for two fewer missed spam messages.")
    p("Several alternatives have higher test F1 under particular decision rules. They do not replace the CV choice. Tuning reduces test F1 for several variants, including the character/gentle model: its recall increases, but 19 of 904 legitimate messages become false positives (2.10% FPR). This illustrates the difference between the validation objective and a guarantee on unseen messages. All eleven models are retained to make these tradeoffs visible.")
    h("6 Conditional uncertainty")
    p("For the fixed selected model, we sample whole test groups with replacement separately within ham and spam, preserving all messages in each selected group. We use 5,000 draws, seed 42, and the 2.5th and 97.5th percentiles of the resulting metric distribution [8]. The default protocol contains 1,032 test groups. The tuned F1 interval is 0.9249 to 0.9769; recall is 0.8931 to 0.9769 and FPR is 0.11% to 0.88%. Predictions reconstructed from the complete error ledger must match the saved confusion counts before resampling.")
    p("These intervals condition on the fitted model, decision rule, split and observed class/group composition. They exclude training and model-search variability, unknown campaign dependence and domain shift. With few errors, percentile intervals can be discrete or degenerate. They are approximate descriptive uncertainty estimates, not evidence that the same coverage holds in a new SMS population.")

    page()
    h("7 Character obfuscation experiments")
    p("Three nonadaptive scenarios modify raw spam before the fitted preprocessing pipeline. Leet replaces eligible ASCII letters using a fixed substitution map. Separator insertion places a period, hyphen or underscore between adjacent ASCII letters. Unicode substitution replaces eligible Latin letters with specified Cyrillic lookalikes. Intensity is the edit probability per eligible character or boundary, rather than the fraction of all characters changed. We use intensities 0.0 to 0.5 in steps of 0.1 and twenty seeds, 10042 to 10061.")
    p("Every model receives the same attacked messages for each scenario, intensity and seed. Ham remains unchanged, and clean validation thresholds stay fixed. Attack success rate (ASR) is the fraction of initially correctly detected spam that becomes ham after editing; its denominator excludes spam already missed on clean input. Twenty seeds measure randomness of these edits on the same test messages, not uncertainty across new training or test sets.")
    doc.add_picture(str(ROOT / "paper/figures/robustness.png"), width=Inches(7))
    p("Figure 1. Default protocol, gentle cleanup and validation thresholds. Curves show mean spam recall; bands show one attack-seed standard deviation. All classifiers are LinearSVC. Different clean false-positive rates must be considered alongside recall.", "Caption")
    robustness = read_csv("full", "robustness_summary.csv")
    rows = []
    for model, short in [(selected, "Word"), ("linear_svc_char_gentle", "Character"), ("linear_svc_word_char_gentle", "Word + character")]:
        row = next(r for r in clean if r["model"] == model and r["threshold_mode"] == "validation")
        vals = [next(r for r in robustness if r["model"] == model and r["threshold_mode"] == "validation" and r["attack_type"] == attack and float(r["intensity"]) == .3) for attack in ["leet", "separators", "unicode"]]
        rows.append([short, f'{100*float(row["fpr"]):.2f}%', *[f'{float(v["recall_mean"]):.3f} / {float(v["asr_mean"]):.3f}' for v in vals]])
    table(["Features", "Clean FPR", "Leet R / ASR", "Separators R / ASR", "Unicode R / ASR"], rows,
          [1.2, .85, 1.55, 1.75, 1.65], "Table 3. Mean recall (R) and ASR at intensity 0.3, with gentle cleanup and tuned thresholds.")
    p("The selected word model is most sensitive to separator insertion at intensity 0.3: mean recall drops to 0.6035 and ASR reaches 0.3582. Character features retain more recall across these scenarios, but their tuned clean FPR is 2.10%, compared with 0.44% for words. The combined model offers an intermediate measured tradeoff. These observations do not isolate a universal representation advantage: configurations and operating points differ. Human readability, meaning preservation, adaptive optimization and modifications to ham were not evaluated.")

    page()
    h("8 Sensitivity to message templates")
    p("Case and whitespace grouping can separate messages that differ only in a phone number or URL. A diagnostic rule masks URLs and digit runs in normalized messages longer than thirty characters. Short messages retain their normalized key to reduce accidental merging of unrelated short replies. The rule is stateless, uses no model scores and changes grouping only; classifiers still receive original strings. Conflicting labels within a template group are rejected.")
    p("The default split contains thirteen templates crossing partition boundaries, covering thirty-one messages; six test messages have a training counterpart under this rule. A separate complete experiment uses the 5,137 template groups for both holdout splitting and training CV. It produces 3,095 training, 1,040 validation and 1,036 test messages, including 904 ham and 132 spam in test. No template crosses partitions under the declared rule. This does not identify every semantic paraphrase or campaign.")
    template_clean = read_csv("template", "clean_metrics.csv")
    template_ci = read_csv("template", "confidence_intervals.csv")
    rows = []
    for mode, label in [("default", "Native default"), ("validation", "Validation tuned")]:
        r = next(r for r in template_clean if r["model"] == template["selected_model"] and r["threshold_mode"] == mode)
        ci = next(r for r in template_ci if r["metric"] == "f1" and r["threshold_mode"] == mode)
        rows.append([label, f'{float(r["precision"]):.4f}', f'{float(r["recall"]):.4f}', f'{float(r["f1"]):.4f}', f'{r["fp"]}/{r["fn"]}', f'{float(ci["lower"]):.4f} to {float(ci["upper"]):.4f}'])
    table(["Decision rule", "Precision", "Recall", "F1", "FP/FN", "95% F1 interval"], rows,
          [1.3, .8, .75, .75, .7, 2.7], "Table 4. Template protocol, CV-selected character LinearSVC without punctuation.")
    p("Training CV selects the character model without punctuation in this experiment. Tuning raises recall slightly but increases false positives from two to eleven and lowers F1 from 0.9531 to 0.9288. Test FPR reaches 1.22% despite meeting the validation constraint. The conditional intervals use 1,028 template test groups. These findings retain the unfavorable outcome rather than choosing a more flattering threshold after evaluation.")
    p("The template analysis was introduced after inspecting the default experiment and is exploratory. The two protocols change the split, fitted models and selected configurations together. Their difference is not a causal estimate of leakage, and their F1 intervals are not a paired significance test. The default model remains the frozen choice for its original protocol.")
    h("9 Scope of the historical experiments")
    p("Both historical protocols use one English corpus. Row order is not a timestamp, and the available labels do not establish independent senders or campaigns [1]. Related messages may remain after either grouping rule. Section 10 tests the original model on a separate reported-SMS collection. Russian or Kazakh accuracy, performance by language or channel, and URL safety remain unmeasured. A ham prediction does not establish that a link or sender is safe.")
    p("The IMC sample has no legitimate messages. Section 11 adds two-class follow-up evidence, with limited ham coverage. Representative modern ham, reliable source/date annotations and new final data remain necessary for further development. Deployment also requires measurements of latency, memory, privacy behavior and false-positive cost; sparse features alone do not establish on-device suitability.")

    page()
    h("10 External evaluation on reported SMS")
    p("We use the CC BY 4.0 author artifact from Fishing for Smishing [9], pinned at commit a6175560b57387199871e51fbef6bc523d2516b4. It contains 33,869 pseudonymized reports with publisher-supplied text and annotations, partly extracted using language models. We use original English text only. A 2025 release does not imply 2025 message dates. Requiring one explicit 2023 or 2024 year yields 488 rows, mostly 2023; normalized deduplication leaves 372 texts in 371 template groups. The supplementary cohort contains 15,531 distinct English texts across all/unknown years. Neither cohort matches the entire historical corpus under exact, normalized or template keys; paraphrase and campaign independence are not established.")
    p("Source, membership, artifact hash and decision rules were fixed before external prediction (repository PR #8). No fitting or threshold selection occurs. All inputs, including OOV messages, remain in the denominator. Table 5 uses the published pseudonymized text; the two overlapping cohorts are not independent replications.")
    rows = []
    for cohort, label in [("recent_2023_2024", "2023–2024"), ("all_english", "All English")]:
        for mode, rule in [("validation", "Validation"), ("default", "Native")]:
            r = ext(cohort, mode)
            rows.append([label, rule, f'{int(r["detected"]):,} / {int(r["n_messages"]):,}',
                         f'{int(r["missed"]):,}', f'{100*float(r["recall"]):.2f}%',
                         f'{100*float(r["recall_lower"]):.2f}–{100*float(r["recall_upper"]):.2f}%'])
    table(["Cohort", "Decision", "Detected / total", "Missed", "Recall", "95% interval"], rows,
          [1, .9, 1.45, .7, .8, 2.15], "Table 5. Frozen model on positive reports. No external F1, precision, accuracy or FPR estimate.")
    doc.add_picture(str(ROOT / "results/external/imc25/evaluation/recall.png"), width=Inches(7))
    p("Figure 2. Same model and validation threshold. Intervals condition on observed groups; external cohorts overlap. Source, sampling, message type and pseudonymization differ, so the contrast is not a causal estimate of temporal degradation.", "Caption")
    stripped = ext("recent_2023_2024", view="placeholders_removed")
    p(f'The primary result is {primary_recall} recall, with 196 misses. Removing the fixed publisher placeholder tokens gives {100*float(stripped["recall"]):.2f}% (175 detected), a 0.27 percentage-point change; it cannot recover unredacted originals. Intervals use 5,000 whole-template resamples, seed 42, and exclude training, annotation and sampling uncertainty. Category counts are descriptive. This positive-report convenience sample exposes limited transfer but cannot measure the cost of blocking legitimate SMS.')

    page()
    h("11 Adaptation with ordinary messages")
    p("A follow-up treats the inspected IMC reports as development material. We add Ahmed and Islam's CC BY 4.0 financial-message dataset, published in October 2025 [10]. It mixes SMS, chats and email, with no receipt dates or per-message channel. Of 523 rows, a Bengali-script exclusion leaves 422; three duplicates and seven historical overlaps leave 412 before splitting. This is a Latin-script auxiliary sample, not verified English-only modern SMS. The independent Spam of the Times source [11] contains dated SMS from one Philippine recipient. We retain unsolicited spam and government public service announcements (PSAs), exclude ambiguous advertising and redacted private texts, and remove repeats and known near-copies. Its 754 test messages span September 2022 to July 2026, mix languages and include only 55 retained ham messages.")
    p("Source hashes, membership and selection rules were committed before fitting (PR #11). Financial groups use 60/20/20 train/validation/test proportions; IMC uses 80/20. Campaign keys mask entities and numbers; fixed character 5-gram hashing removes near-copies at cosine similarity >=0.85. Historical test groups are excluded from development. After exclusions, training has 15,172 rows, including 99 new ordinary financial messages. Validation has 3,034 rows. The financial test has 32 ham and 46 scam messages; the independent SMS test has 55 PSAs and 699 spam messages. These screens do not certify semantic independence.")
    p("We refit six LinearSVC pipelines: word 1–2 grams, character 3–5 grams and their union, each with C=0.5 or 1. Class totals are equally weighted, then sources and campaigns within each class. Validation maximizes mean source recall subject to FPR <=1% separately on historical and financial ham. It selects word plus character features, C=1 and threshold 0.64907808395485. No refit or tuning follows test inspection.")
    rows = []
    for source, label in [("financial", "Financial"), ("sms_history", "SMS history")]:
        for model, model_label in [("baseline", "Original"), ("candidate", "Adapted")]:
            r = next(r for r in adapted if r["source"] == source and r["model"] == model)
            rows.append([label, model_label, f'{100*float(r["recall"]):.2f}%', f'{float(r["f1"]):.4f}',
                         f'{r["fp"]} / {r["n_ham"]}', f'{r["fn"]} / {r["n_spam"]}'])
    table(["Test source", "Model", "Recall", "Spam F1", "FP / ham", "FN / spam"], rows,
          [1.2, 1, 1, 1, 1.4, 1.4], "Table 6. Adaptation improves spam detection but increases errors on independent PSAs.")
    p("The independent source gains 112 spam detections but eight additional false positives. Paired conditional 95% group-bootstrap intervals are +0.0998 to +0.1721 for F1 change and +1.79 to +28.07 percentage points for FPR change (5,000 resamples, seed 20261018). ROC-AUC falls from 0.6979 to 0.6466. The financial holdout's F1 change interval is +0.1102 to +0.3457; its FPR-change interval includes zero. The predeclared replacement criterion requires a positive 95% lower bound for independent F1 improvement and FPR no worse than baseline and <=1% on each new test source. It fails. The candidate remains experimental and the default is unchanged. Small, narrow ham samples and source/language differences limit generalization.")

    page()
    h("12 Reproducibility and conclusion")
    p("The repository supplies split/CV manifests, all 84 search results, validation and test metrics, attack realizations, error ledgers, intervals and fitted-pipeline export. Dataset, source code and artifacts have SHA-256 records. Both historical protocols were independently verified, including all 7,920 attack metric rows per run. Python 3.12.14, scikit-learn 1.9.1 and exact dependencies are recorded. scripts/verify_external.py reconstructs the separate IMC scores, decisions and summaries from the original fitted pipeline. Verification requires hash-matched local models; binaries are excluded from Git and runs cannot silently overwrite frozen outputs.")
    p("The adaptation follow-up records six fitted candidates, 18,204 validation scores and 1,664 test scores. scripts/verify_adaptation.py independently checks thresholds by exhaustive search, metrics and paired intervals by resampling message rows through whole groups. No model fitting occurs during verification. docs/MODERN_ADAPTATION.md documents the sources, fixed protocol, failed replacement criterion and reproduction commands.")
    p("The word model combines high historical F1 with sensitivity to separator edits and limited source transfer. Refitting on newer messages improves spam recall and same-source financial F1, but increases false positives on independent government notifications. This measured tradeoff prevents replacing the default. Further progress needs more diverse genuine ordinary SMS and a new untouched test; the reported holdouts must not be reused as independent evidence after tuning on these findings.")
    p("Code and experiment artifacts: https://github.com/altynbk/sms-spam-detection")
    h("References")
    refs = [
        "[1] T. Almeida and J. Hidalgo. SMS Spam Collection. UCI Machine Learning Repository, 2012. Dataset DOI: 10.24432/C5CC84. https://archive.ics.uci.edu/dataset/228/sms+spam+collection",
        "[2] T. A. Almeida, J. M. G. Hidalgo and A. Yamakami. Contributions to the study of SMS spam filtering: new collection and results. ACM Symposium on Document Engineering, 2011. Publication record linked from [1].",
        "[3] M. A. Uddin et al. ExplainableDetector: Exploring Transformer-based Language Modeling Approach for SMS Spam Detection with Explainability Analysis. arXiv:2405.08026, 2024. https://arxiv.org/abs/2405.08026",
        "[4] Y. Li, R. Zhang, W. Rong and X. Mi. SpamDam: Towards Privacy-Preserving and Adversary-Resistant SMS Spam Detection. arXiv:2404.09481, 2024. https://arxiv.org/abs/2404.09481",
        "[5] Harichandana B S S et al. COPS: A Compact On-device Pipeline for real-time Smishing detection. IEEE CCNC, 2024. arXiv:2402.04173. https://arxiv.org/abs/2402.04173",
        "[6] F. Pedregosa et al. Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, vol. 12, pp. 2825-2830, 2011. https://jmlr.org/papers/v12/pedregosa11a.html",
        "[7] Scikit-learn developers. Tuning the decision threshold for class prediction. User Guide, accessed 9 October 2026. https://scikit-learn.org/stable/modules/classification_threshold.html",
        "[8] SciPy developers. scipy.stats.bootstrap. Reference documentation, accessed 9 October 2026. https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.bootstrap.html",
        "[9] S. Agarwal, A. Papasavva, G. Suarez-Tangil and M. Vasek. Fishing for Smishing: Understanding SMS Phishing Infrastructure and Strategies by Mining Public User Reports. ACM IMC, 2025. DOI: 10.1145/3730567.3764431. Dataset: https://github.com/reportsmishing/Smishing-Dataset-IMC25",
        "[10] A. R. Ahmed and G. F. Islam. Financial scams detection dataset. Mendeley Data, version 2, 22 October 2025. CC BY 4.0. DOI: 10.17632/znsk27yk3h.2. https://data.mendeley.com/datasets/znsk27yk3h/2",
        "[11] S. L. Chua. Spam of the Times: spam and marketing SMS. CC BY 4.0. Snapshot ad50c348c53f738645f78ee8baa8be5082a1de91, 20 August 2026. https://github.com/scottleechua/data/tree/ad50c348c53f738645f78ee8baa8be5082a1de91/spam-and-marketing-sms",
    ]
    for ref in refs:
        para = p(ref)
        para.paragraph_format.space_after = Pt(7)
        for run in para.runs:
            run.font.size = Pt(9.5)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
